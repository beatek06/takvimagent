"""Yerel bir .xlsx dosyasından SADECE TAKVIM sekmesini test verisi olarak çıkarır.

Kullanım: python dev/extract_fixture.py "<xlsx yolu>"
Çıktı: dev/fixtures/takvim.json (gitignore'da; GitHub'a gitmez)
"""
import json
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl

EPOCH = date(1899, 12, 30)


def convert(v):
    if isinstance(v, datetime):
        return (v.date() - EPOCH).days  # Sheets UNFORMATTED_VALUE gibi seri sayı
    return v


wb = openpyxl.load_workbook(sys.argv[1])
ws = wb["TAKVIM"]
grid = [
    [convert(v) for v in row]
    for row in ws.iter_rows(
        min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column, values_only=True
    )
]
out = Path(__file__).parent / "fixtures" / "takvim.json"
out.parent.mkdir(exist_ok=True)
out.write_text(json.dumps(grid, ensure_ascii=False), encoding="utf-8")
print(f"{len(grid)} satır yazıldı -> {out}")
